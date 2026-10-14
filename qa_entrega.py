"""Construye registros derivados; no escribe en las fuentes oficiales."""
from pathlib import Path
from datetime import datetime, timezone, timedelta
import csv,json,hashlib,subprocess,shutil,os,xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent;OUT=Path(os.environ.get('SEDAPAL_ENTREGA_DIR',str(ROOT/'ENTREGA_PARA_APROBACION')));EV=OUT/'EVIDENCIAS';DL=Path('C:/Users/User/Downloads')
STAMP=datetime.now(timezone(timedelta(hours=-5))).isoformat()
def git(*args):return subprocess.check_output(['git',*args],cwd=ROOT,text=True,encoding='utf-8').strip()
HEAD=git('rev-parse','HEAD');BASE=git('rev-parse','41070eb');VERSION='0.1.1-SNAPSHOT'
def write(name,text):(OUT/name).write_text(text,encoding='utf-8')
def table(name,rows):
 rows=list(rows)
 with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def junit(path):
 rows=[]
 for f in sorted(path.glob('TEST-*.xml')):
  e=ET.parse(f).getroot()
  for t in e.findall('testcase'):
   result='FAIL' if t.find('failure') is not None or t.find('error') is not None else ('NO EJECUTADA' if t.find('skipped') is not None else 'PASS')
   fecha=e.attrib.get('timestamp') or datetime.fromtimestamp(f.stat().st_mtime,timezone(timedelta(hours=-5))).isoformat()+' (última escritura del reporte de esta ejecución)'
   rows.append(dict(id=t.attrib['classname'].split('.')[-1]+'.'+t.attrib['name'],fecha=fecha,duracion_s=t.attrib.get('time',''),estado=result,evidencia=str(f.relative_to(OUT)).replace('\\','/')))
 return rows
src=Document(DL/'SEDAPAL_Catalogo_RF_Backlog_RNF_DoD_v1.4_2026-10-14.docx')
web=json.loads((EV/'navegador-resultados.json').read_text(encoding='utf-8'))
perf=json.loads((EV/'rendimiento-20.json').read_text(encoding='utf-8'))
java=junit(EV/'junit-final');baseline=junit(EV/'junit-baseline')
assert len(java)>0 and all(r['estado']=='PASS' for r in java)
wf=[r for r in web if not r['id'].startswith('BASE-')];wb=[r for r in web if r['id'].startswith('BASE-')]
assert len(wb)==17
assert len(wf)==41 and all(r['estado']=='PASS' for r in wf), 'No generar un informe conforme con una ejecución incompleta o fallida del incremento'
sources=[]
for name in ['ENCARGO_PARA_CODEX_QA_SEDAPAL_2026-10-14.md','SEDAPAL_Catalogo_RF_Backlog_RNF_DoD_v1.4_2026-10-14.docx','SEDAPAL_Matriz_Maestra_Trazabilidad_v1.8_2026-10-14.xlsx','REPORTE_QA_Automatizada_SEDAPAL_v1.0_2026-10-14.docx','SEDAPAL_MVP_Git_Historial_v0.1.0_2026-10-14.bundle','SEDAPAL_MVP_Demo_Java21_v0.1.0_2026-10-14.jar']:
 f=DL/name;sources.append(dict(archivo=str(f),bytes=f.stat().st_size,sha256=hashlib.sha256(f.read_bytes()).hexdigest(),fecha_verificacion=STAMP,uso='Solo lectura'))
table('FUENTES_Y_SHA256.csv',sources)
comparacion=[]
for name in git('ls-tree','-r','--name-only',BASE).splitlines():
 f=ROOT.parent/name
 if f.is_file():
  old=subprocess.check_output(['git','show',f'{BASE}:{name}'],cwd=ROOT)
  comparacion.append(dict(archivo=name,coincide_con_bundle_contenido_normalizado=old.replace(b'\r\n',b'\n')==f.read_bytes().replace(b'\r\n',b'\n')))
table('COMPARACION_CARPETA_ORIGINAL_BUNDLE.csv',comparacion)
HU_RF={1:'RF-01;RF-02',2:'RF-03;RF-04',3:'RF-05',4:'RF-06',5:'RF-07',6:'RF-08',7:'RF-09',8:'RF-10'}
CA_TEST={
 1:['WEB-06;HU01RechazaDatosObligatoriosVacios','HU01RechazaDatosObligatoriosVacios;WEB-06','WEB-06','WEB-06','WEB-06','HU08AltaHashTitularPersistenteSinSuplantacion;WEB-06'],
 2:['WEB-01','WEB-01;HU02ConsultaPublicaNoExponeIdentidadNiObservaciones','WEB-01','WEB-01;HU05HU06CierreExplicitoReaperturaHistorialYNotificacion','WEB-02'],
 3:['WEB-13','WEB-11;WEB-13','WEB-12;WEB-13;WEB-14','WEB-11'],
 4:['WEB-15;registraReclamoConCodigoHistorialYNotificacion','HU04CadaTransicionNotificaCodigoDestinatarioYFecha;HU04AvisosNoSeMezclan;WEB-15','WEB-15','HU04CadaTransicionNotificaCodigoDestinatarioYFecha;HU04AvisosNoSeMezclan'],
 5:['WEB-16;WEB-17;HU05RechazaCierreNoAtendido;HU05HU06AjenoYOperadorNoPuedenCerrar','WEB-16;WEB-17','HU05HU06CierreExplicitoReaperturaHistorialYNotificacion','WEB-17'],
 6:['WEB-19;WEB-20','HU06AceptaLimitesMotivo;HU06RechazaFueraDeLimites;WEB-18','WEB-19','WEB-19','WEB-19;HU05HU06CierreExplicitoReaperturaHistorialYNotificacion'],
 7:['WEB-10','ArchivoServiceTest;RNF06RechazaExtensionOContenidoDiscrepante;RNF06RechazaExcesoTamanoYConteo;RNF06RechazaPdfConFirmaPeroSinEstructura;RNF06AceptaPdfCifradoReconocidoSinPublicarlo;UPLOAD-01;UPLOAD-02;UPLOAD-03','WEB-10;HU03HU07RevisionYDescargaExigeTitular','WEB-11;WEB-13'],
 8:['WEB-03','HU08RechazaCadaDatoObligatorioAusente;HU08RechazaClaveUnicodeExcedidaEnBytes;HU08RechazaClaveInsuficiente;WEB-04','WEB-05;BASE-HTTP-09','WEB-05;WEB-06','WEB-09;WEB-12;WEB-14;WEB-20','WEB-07','WEB-08']}
cas=[]
for n,idx in [(1,2),(2,3),(3,4),(4,5),(5,6),(6,7),(7,8),(8,16)]:
 for k,row in enumerate(src.tables[idx].rows[1:]):
  ident,text=[c.text for c in row.cells];limit='Cobertura del escenario documentado; no certifica DoD.'
  estado='Probado con éxito en escenarios ejecutados'
  if (n,k)==(7,1):limit='Tipos y límites probados en servicio; PDF cifrado reconocido sin publicación automática; negativos multipart HTTP comprobados. Sexto adjunto probado en servicio, no mediante UI.'


  cas.append(dict(HU=f'HU-{n:02}',CA=ident,RF=HU_RF[n],criterio_aprobado=text,estado=estado,pruebas=CA_TEST[n][k],evidencia='EVIDENCIAS/navegador-resultados.json;EVIDENCIAS/junit-final',limites=limit,fecha=STAMP,version_commit=HEAD,responsable='Codex agente automatizado'))
assert len(cas)==39
table('ESTADO_39_CA.csv',cas)
hus=[]
for n in range(1,9):
 matching=[c for c in cas if c['HU']==f'HU-{n:02}']
 hus.append(dict(HU=f'HU-{n:02}',RF=HU_RF[n],implementacion='Implementado en incremento local; aceptación pendiente',pruebas='Probado con éxito en escenarios ejecutados; ver ESTADO_39_CA.csv',CA_total=len(matching),CA_con_cobertura_parcial=sum(c['estado']=='Parcial' for c in matching),DoD='NO DONE',pendiente='RNF-03, RNF-09 completos; DoD-08 aprobación e incorporación oficial; revisión del equipo',antecedente='Código heredado parcial' if n in [1,2,4,7,8] else 'Brecha funcional en antecedente',evidencia='ESTADO_39_CA.csv;EVIDENCIAS/navegador-resultados.json',version_commit=HEAD))
table('ESTADO_8_HU.csv',hus)
rfrows=[]
for row in src.tables[0].rows[1:]:
 code,title,definition,approval=[c.text for c in row.cells];hu=next(n for n,refs in HU_RF.items() if code in refs.split(';'))
 rfrows.append(dict(RF=code,descripcion_aprobada=definition,HU=f'HU-{hu:02}',implementacion=hus[hu-1]['implementacion'],resultado='Cobertura técnica parcial; no aceptación final',evidencia='ESTADO_39_CA.csv',version_commit=HEAD))
table('ESTADO_10_RF.csv',rfrows)
RNF=[
 ('Parcial','12 comprobaciones automáticas de overflow y capturas Chrome a 1366/768/360; cierre/reapertura capturados. Falta cubrir todas las vistas internas y revisión completa de dispositivos.','RESP-*;capturas PNG'),
 ('Parcial','Chrome y Edge headless en Windows. Safari iOS y dispositivos físicos no disponibles; tamaños emulados no equivalen a un móvil real.','EDGE-01;chrome-version.txt;edge-version.txt'),
 ('NO EJECUTADA','No hay cinco participantes reales disponibles; protocolo preparado. No se sustituyen tiempos humanos por latencia.','PROTOCOLO_USABILIDAD.md'),
 ('PASS escenario local',f"{perf['cumplen']}/20 consultas DOM <=3 s; p95={perf['p95']:.3f} s. {perf['escenario']}. No extrapolar a producción.",'rendimiento-20.json'),
 ('Probado con éxito en escenarios ejecutados','Autorregistro, titularidad, privacidad de evidencia, CSRF, permisos y sesión. Sin pentest exhaustivo ni revisión independiente.','WEB-03 a WEB-20;IncrementoIntegrationTest'),
 ('Probado con éxito en escenarios ejecutados','JPG/PNG decodificados; PDF analizado con PDFBox; límite 5 MiB, sexto adjunto y PDF cifrado probados en servicio; formato, contenido y exceso rechazados por HTTP.','ArchivoServiceTest;RNF06*;WEB-10;UPLOAD-01 a UPLOAD-03'),
 ('Parcial','BCrypt comprobado y accesos ajenos denegados. Pruebas por HTTP de loopback; HTTPS desplegado no ejecutado. Credenciales demo solo para evaluación local.','HU08AltaHashTitularPersistenteSinSuplantacion;WEB-14;WEB-20'),
 ('Probado con éxito en escenarios ejecutados','Historial anterior preservado en cierre/reapertura. Registro de titular/fecha y revisión de evidencia persistidos. No se probó restauración ante pérdida de almacenamiento.','WEB-19;HU05HU06CierreExplicitoReaperturaHistorialYNotificacion'),
 ('Parcial','Tab/foco y recorrido de consulta, cierre y reapertura con teclado; 59 muestras sólidas de la pantalla de seguimiento con contraste >=4.5:1; no representa todas las pantallas. Falta contraste íntegro de todas las vistas y recorrido de cada función.','A11Y-01;KEYBOARD-01;contraste-calculo.csv;contraste-muestras.json')]
rnfrows=[]
for i,row in enumerate(src.tables[12].rows[1:]):
 code,category,definition=[c.text for c in row.cells];state,detail,evidence=RNF[i]
 rnfrows.append(dict(RNF=code,categoria=category,requisito_aprobado=definition,estado=state,resultado_o_causa=detail,evidencia=evidence,fecha=STAMP,version_commit=HEAD))
table('ESTADO_9_RNF.csv',rnfrows)
dod=[]
states=[('Conforme documental','Catálogo aprobado por declaración de la fuente; conformidad original por archivar.'),('Parcial','Implementado y ejecutado localmente; limitaciones en CA/RNF registradas.'),('Parcial','39 CA revisados; 39 con correspondencia a pruebas satisfactorias en escenarios ejecutados; validación formal de casos pendiente.'),('Pendiente','RNF-03, RNF-09 y compatibilidad/HTTPS incompletos.'),('Parcial','Auditoría automatizada del mismo agente; revisión independiente por equipo pendiente.'),('Conforme técnico','Historial y commits reales recuperables por bundle; versión 0.1.1-SNAPSHOT.'),('Conforme técnico','JUnit XML, JSON, logs y capturas fechadas; responsable automatizado identificado.'),('Pendiente aprobación','CSV derivado preparado; matriz oficial no modificada ni aprobada automáticamente.'),('Parcial','Incremento disponible para inspección; no consta revisión ni aceptación humana.')]
for i,row in enumerate(src.tables[17].rows[1:]):
 code,cond,evidence=[c.text for c in row.cells];state,reason=states[i]
 dod.append(dict(DoD=code,condicion_aprobada=cond,estado=state,detalle=reason,aplicacion='HU-01 a HU-08; ninguna Done',evidencia='REPORTE_QA.md;ESTADO_39_CA.csv',version_commit=HEAD))
table('ESTADO_9_DOD.csv',dod)
nc=[
 ('NC-01','Alta','HU-08','Faltaba autorregistro y primera incidencia de cuenta nueva','Corregida y probada','WEB-03 a WEB-08;HU08AltaHashTitularPersistenteSinSuplantacion','Alta ciudadana con rol fijo y referencia local si no tiene suministro'),
 ('NC-02','Alta','HU-03;HU-07','No existía consulta privada ni aprobación de evidencias','Corregida y probada','WEB-10 a WEB-14','Descarga por titular; operador revisa antes de publicación'),
 ('NC-03','Alta','HU-05','Cierre por operador en lugar de confirmación del titular','Corregida y probada','WEB-16;WEB-17;HU05HU06AjenoYOperadorNoPuedenCerrar','Cierre reservado al titular con POST y confirmación expresa'),
 ('NC-04','Alta','HU-06','No existía solicitud de reapertura','Corregida y probada','WEB-18 a WEB-20;HU06AceptaLimitesMotivo','Motivo 10-500; nueva transición y preservación del historial'),
 ('NC-05','Alta','RNF-05','POST ajeno devolvía 405 al reenviar a vista GET de acceso denegado','Corregida; regresión navegador','EVIDENCIAS/INTENTO_01/navegador-resultados.json;WEB-20','Handler devuelve 403 sin reenvío de método POST'),
 ('NC-06','Alta','RNF-06','Se confiaba en MIME; faltaban tipo real y cupo','Corregida en escenarios ejecutados','ArchivoServiceTest;RNF06*','Decodificar imágenes; analizar PDF; limitar tamaño y conteo'),
 ('NC-07','Media','RNF-03','Falta prueba con participantes reales','Abierta','PROTOCOLO_USABILIDAD.md','Equipo organiza cinco participantes; no se fijan citas ni resultados ficticios'),
 ('NC-08','Media','RNF-09','Accesibilidad completa y contraste aún sin verificar','Abierta','A11Y-01;KEYBOARD-01','Recorrido íntegro y medición de contraste por pantalla'),
 ('NC-09','Media','RNF-02;RNF-01','Safari iOS, dispositivos reales y vistas internas incompletas','Abierta','EDGE-01;RESP-*','Obtener acceso a entornos y ampliar matriz de navegador'),
 ('NC-10','Media','DoD-08;CM;PPQA','Matriz oficial y conformidades pendientes de incorporación; revisión independiente no realizada','Abierta','ACTUALIZACIONES_PARA_MATRIZ.csv','Responsable valida propuestas; revisor del equipo verifica incremento'),
 ('NC-11','Media','RNF-07','HTTPS de entorno desplegado no comprobado','Abierta','entorno-final.json','Validar en entorno HTTPS autorizado; no publicar desde esta auditoría'),
 ('NC-12','Media','HU-07/CA-02;RNF-06','La validación inicial rechazaba PDF cifrado','Corregida y probada','RNF06AceptaPdfCifradoReconocidoSinPublicarlo;UPLOAD-01;UPLOAD-02;UPLOAD-03','Parser reconoce cifrado; conserva bytes y exige revisión antes de publicar'),
 ('NC-13','Baja','HU-04;HU-08','Faltaban aserciones por evento y campo obligatorio','Corregida y probada','HU04CadaTransicionNotificaCodigoDestinatarioYFecha;HU08RechazaCadaDatoObligatorioAusente','Se añadieron pruebas por eventos, campos y contraseña Unicode'),
 ('NC-14','Media','RNF-06','Rollback de BD después de guardar adjunto podía dejar archivo huérfano','Corregida y probada','ArchivoRollbackIntegrationTest.eliminaSoloArchivoNuevoCuandoRevierteLaBaseDeDatos','Sincronización transaccional elimina únicamente el archivo nuevo al revertir'),
]
nc.extend([('NC-15','Media','RNF-05','El JAR original devuelve 405 al omitir CSRF autenticado','Corregida en incremento; original conservado','BASE-HTTP-03;SECURITY-CSRF-01','Mantener originales intactos y usar el incremento con respuesta 403'),('NC-16','Baja','QA automatizada','El verificador consultaba antes de completar los datos demo','Corregida en verificador','EVIDENCIAS/INTENTO_03;qa_navegador.ready','Esperar caso demo Asignado antes de los checks; preservar los resultados anteriores')])
table('NO_CONFORMIDADES.csv',[dict(id=a,prioridad=b,referencias=c,hallazgo=d,estado=e,evidencia=f,accion_propuesta=g,responsable_seguimiento='Equipo 5 por asignar; auditoría Codex',fecha=STAMP) for a,b,c,d,e,f,g in nc])
testrows=[]
for r in java:
 refs=';'.join(c['HU']+'/'+c['CA'] for c in cas if any(x in r['id'] for x in c['pruebas'].split(';')))
 testrows.append(dict(id=r['id'],referencias=refs or 'Regresión técnica/RNF;ver clase Java',fecha=r['fecha']+' (reloj entorno Lima)',entorno='Java 21.0.9;Maven 3.9.15;Spring test;H2 memoria',version_commit=HEAD,responsable='Codex agente automatizado',entradas='Datos ficticios definidos en el método Java; ver CODIGO_FUENTE/src/test/java',esperado='Se cumplen todas las aserciones identificadas en el método',obtenido=r['estado'],estado=r['estado'],duracion_s=r['duracion_s'],evidencia=r['evidencia'],pasos='mvn -DforkCount=0 clean verify; localizar método por ID'))
for r in web:
 testrows.append(dict(id=r['id'],referencias=r['referencias'],fecha=r['fecha'],entorno='Windows;Chrome/Edge headless;demo/H2;loopback;ver entorno JSON',version_commit=r['version'],responsable=r['responsable'],entradas='Cuentas y casos ficticios del perfil demo; qaCiudadano; archivo evidencia-ficticia.pdf;ver qa_navegador.py',esperado=r['esperado'],obtenido=str(r['obtenido']),estado=r['estado'],duracion_s=r['duracion_s'],evidencia=r['evidencia'],pasos=r['pasos']))
table('RESULTADOS_PRUEBAS.csv',testrows)
edt={1:'1.2.1;1.2.2;1.2.3',2:'1.3.1;1.3.2;1.3.3',3:'1.4.2',4:'1.4.3',5:'1.5.1',6:'1.5.2',7:'1.4.1',8:'CF-001 (sin EDT aprobado)'}
table('ACTUALIZACIONES_PARA_MATRIZ.csv',[dict(EDT=edt[int(c['HU'][-2:])],RF=c['RF'],HU=c['HU'],CA=c['CA'],implementacion='CODIGO_FUENTE/src;ver CAMBIOS_IMPLEMENTADOS.md',evidencia=c['evidencia'],prueba=c['pruebas'],PPQA='Auditoría automatizada;observaciones en NO_CONFORMIDADES.csv',version_commit=HEAD,estado_propuesto=c['estado']+';NO DONE',aprobacion='PENDIENTE;no modifica matriz oficial') for c in cas])
history=git('log','--date=iso-strict','--format=%H | %aI | %an | %s')
write('CAMBIOS_IMPLEMENTADOS.md',f'''# Cambios implementados y procedencia

Fecha de corte {STAMP}. Versión {VERSION}. Commit del código probado `{HEAD}`.

El bundle contiene la importación antecedente `fc52549` y la mejora HU-02 `41070eb`. Se recuperó el historial completo; no se atribuye este trabajo a Sprints 1 o 2. El POM heredado declaraba 0.0.1-SNAPSHOT pese a la denominación del paquete v0.1.0. El incremento usa 0.1.1-SNAPSHOT.

La carpeta original se contrastó contra el árbol de 41070eb; resultados en COMPARACION_CARPETA_ORIGINAL_BUNDLE.csv. Los documentos oficiales se leyeron sin escritura; sus SHA-256 están en FUENTES_Y_SHA256.csv. El ZIP de código mencionado no está en Descargas; se utilizó la carpeta extraída y el bundle disponible, sin bloquear el trabajo.

Cambios del curso ejecutados en esta auditoría: alta ciudadana y hash BCrypt; referencia local de ubicación para primera incidencia; consulta pública con campos permitidos; cierre titular y solicitud pendiente de reapertura; descargas privadas y aprobación por operador; validación de JPG/PNG/PDF y cupo; corrección HTTP 403; generación de código por identidad de BD; pruebas nuevas y registros QA. No hay integración institucional, nueva EDT aprobada ni cambios en RF/HU/RNF/DoD.

## Historial real

```text
{history}
```

El bundle entregado permite recuperar el repositorio. Las pruebas referencian el commit del código ejecutado; un commit posterior de empaquetado solo agrega herramientas y documentación.
''')
write('PROTOCOLO_USABILIDAD.md','''# Protocolo pendiente de usabilidad

Estado NO EJECUTADA. RNF-03 exige participantes reales; esta auditoría no dispone de ellos.

Propuesta para el Equipo 5: seleccionar al menos cinco participantes y obtener su consentimiento; usar un caso ficticio con evidencia publicada. Entregar un código y medir desde el inicio de la búsqueda hasta que la persona identifique estado, última actualización y evidencia. Registrar dispositivo, navegador, versión, ayuda recibida, tiempo individual, éxito y observaciones. Calcular promedio de tiempos válidos y compararlo con 30 segundos; conservar fallos y abandonos. El evaluador no debe guiar durante la tarea. Anonimizar participantes mediante P01 a P05. No se han asignado personas, fechas de sesión ni resultados.
''')
write('PLAN_CORRECTIVO.md','''# Plan correctivo para revisión

Las prioridades describen orden técnico, no un compromiso de Sprint, presupuesto o fecha.

1. Revisar las pruebas de RNF-06 y HU-07, incluidos PDF cifrado, negativos multipart y limpieza en rollback, antes de aceptar el incremento.
2. Revisar las aserciones añadidas por cada evento y campo obligatorio; ampliar casos de estrés si el equipo define un escenario de evaluación adicional.
3. Completar accesibilidad y pantallas internas; obtener Safari iOS/dispositivos reales y HTTPS de un entorno autorizado. No se solicita publicación externa.
4. Ejecutar RNF-03 con cinco participantes reales siguiendo el protocolo; responsabilidad nominal pendiente del Equipo 5.
5. Revisor distinto del implementador revisa PPQA; responsable valida las propuestas CSV e incorpora la matriz oficial, conformidades y evidencia CF-001 en CM.
6. Revisar cronograma/EDT de HU-08 y formalizar Sprint Backlogs mediante los eventos reales del equipo. No atribuir la auditoría actual a Sprints pasados.

No se estima costo, duración ni fecha sin capacidad real y planificación del equipo. Ninguna HU se contabiliza como Done durante estos pendientes.
''')
write('MATRIZ_PRELIMINAR_ESTADO.md',f'''# Estado preliminar de seguimiento y procesos

Corte {STAMP}. Registros derivados para aprobación; no equivalen a evaluación formal CMMI ni a aceptación del equipo.

| Área o artefacto | Avance verificable | Pendiente |
|---|---|---|
| PP y EDT | Correspondencias a EDT existentes extraídas de matriz v1.8 | Validar impacto de HU-08 y cronograma; no inventar estimaciones |
| REQM | Diez RF, ocho HU y 39 CA conservados y revisados | Archivar evidencia de conformidad original y CF-001 |
| PMC | Registro fechado de defectos y plan correctivo | Asignar responsables, fechas y seguimiento por equipo |
| MA | {len(java)} pruebas JUnit y {len(wf)} checks web del incremento; muestras de 20 consultas | No hay horas de equipo, HC aprobadas ni tiempos de participantes |
| PPQA | Auditoría automatizada y revisión estática, intento fallido conservado | Revisión independiente del equipo; NC abiertas |
| CM | Bundle original verificado, rama local y commits reales, hashes de fuentes | Aprobar entrega y actualizar línea base documental |
| SAM | No se proporcionaron contratos ni gestión de proveedores | Aplicabilidad por determinar por el equipo; no se inventan contratos |
| Sprint Backlogs | Sprints previstos en catálogo, sin atribución de trabajo retrospectiva | No consta Planning ni compromiso real; formalización pendiente |
| APF2 | Evidencia técnica y propuestas disponibles como insumo | Informe académico consolidado y aceptación humana pendientes |

DoD: 0/8 HU declaradas Done. HC no puede incrementarse con checks técnicos. Se propone registrar HC verificada=0 para este incremento; cualquier fórmula con historias planificadas requiere un denominador aprobado, no el mero total del Product Backlog. La razón de checks PASS mide únicamente la suite ejecutada y no el cumplimiento del alcance.
''')
write('REGISTRO_MA.md',f'''# Registro de medición y análisis

Corte {STAMP}. Fuente de medición: JUnit XML, JSON Playwright y rendimiento-20.json. Responsable de ejecución: Codex, agente automatizado. Revisión humana pendiente.

- Línea base 41070eb: {len(baseline)} tests JUnit actuales; no se reutilizan reportes antiguos.
- Incremento {HEAD}: {len(java)} tests JUnit, {sum(r['estado']=='PASS' for r in java)} PASS, sin omitidos.
- Humo del JAR original: {len(wb)} checks actuales; ocho HU-02 y nueve HTTP reconstruidos a partir del alcance del informe (no se dispone del script HTTP original).
- Navegador del incremento: {len(wf)} checks; {sum(r['estado']=='PASS' for r in wf)} PASS. No sumar checks solapados como historias.
- Rendimiento: {perf['cumplen']}/20 <=3 s; p95 {perf['p95']:.3f} s, rango {min(perf['segundos']):.3f}–{max(perf['segundos']):.3f} s. Método nearest-rank, 19.ª muestra ordenada. {perf['escenario']}.
- Usabilidad humana: NO EJECUTADA; sin promedio válido. HC según DoD: 0 HU declaradas Done. Productividad, esfuerzo real del equipo y cumplimiento de Sprint: no calculables con las fuentes disponibles.

No se extrapolan las mediciones a producción ni se afirma independencia entre todos los checks.
''')
write('CHECKLIST_PPQA.md',f'''# Registro de revisión PPQA

Fecha {STAMP}; implementador y auditor automatizado Codex. Este registro requiere revisión independiente por el Equipo 5.

Se revisaron las 39 condiciones de aceptación, nueve RNF y nueve condiciones DoD contra el catálogo oficial, sin cambiar su texto. Resultados por ítem en CSV. Se verificó la procedencia Git, se ejecutó compilación limpia, JUnit, navegación y accesos negativos. Se inspeccionaron capturas reales. El intento 01 conserva un fallo HTTP corregido mediante commit 91511b4. Las limitaciones abiertas están en NO_CONFORMIDADES.csv.

Dictamen provisional: incremento disponible para inspección técnica; no aceptar ninguna HU como Done. Falta validación humana, RNF pendientes y actualización oficial de trazabilidad. Sin firma ni aprobación inventada.
''')
write('REGISTRO_CM.md',f'''# Registro de configuración

Línea base documental: catálogo v1.4, matriz v1.8, RF/HU v1.1, RNF y DoD v1.0 según fuentes. Línea base de código: `{BASE}`, bundle completo verificado. Rama de trabajo `auditoria/2026-10-14`. Incremento probado `{HEAD}`, Maven {VERSION}.

Las fuentes oficiales se conservan fuera de esta entrega y no se escribieron. El control de cambios diferencia antecedente fc52549, mejora HU-02 41070eb y commits de auditoría. La entrega incluye fuente exportada, bundle Git, historial y SHA-256 de archivos. Se propone incorporar los CSV a la matriz solo después de la aprobación del responsable. No se crearon tags de liberación aprobada, repositorios remotos ni despliegues externos.

Recuperación: `git clone HISTORIAL_GIT.bundle sedapal-revision`; luego comprobar `git log --all --oneline` y el commit probado. El empaquetado contiene datos sintéticos de pruebas; no incluye bases H2 de trabajo ni archivos de usuario reales.
''')
summary=f'''# Resumen ejecutivo para el Equipo 5

La auditoría recuperó y conservó el historial Git del MVP, incluida la mejora HU-02 41070eb. El incremento {VERSION} implementa cuenta ciudadana, seguimiento público sin información privada, evidencias autorizadas y revisadas, cierre por titular y solicitud de reapertura con historial preservado. No se alteraron los documentos oficiales ni se publicó el sistema.

Resultados actuales: línea base {len(baseline)} tests JUnit; incremento {len(java)} tests JUnit PASS; humo original {sum(r['estado']=='PASS' for r in wb)}/17 checks PASS; incremento {sum(r['estado']=='PASS' for r in wf)}/{len(wf)} checks de navegador PASS. Hay capturas reales a 1366, 768 y 360 px y ejecución en Chrome/Edge. Se corrigió un defecto HTTP 405 detectado en navegador; el intento fallido se conserva. Rendimiento local: {perf['cumplen']}/20 consultas <=3 s, p95 {perf['p95']:.3f} s; muestra pequeña y carga secuencial.

Ninguna HU se declara Done. Siguen pendientes usabilidad con participantes reales, accesibilidad y contraste completos, Safari/dispositivos físicos, HTTPS desplegado. No hay constancia de revisión independiente ni aceptación del equipo. Los 39 CA tienen trazabilidad propuesta y limitaciones explícitas.

Para aprobar: revisar el incremento y las NC abiertas; decidir responsables y planificación real, revisar el tratamiento de archivos y actualizar la matriz oficial a partir de los CSV. Esta entrega no aprueba cambios de alcance, Sprints, presupuesto ni fechas. El ZIP y el bundle permiten inspección y recuperación posterior.
'''
write('RESUMEN_EJECUTIVO.md',summary)
report=f'''# Informe QA del incremento académico SEDAPAL

Corte {STAMP}. Elaborado por Codex, agente automatizado, para el responsable delegado del Equipo 5. Estado: para aprobación. Código probado `{HEAD}`; versión {VERSION}. Este informe documenta ejecución técnica y pendientes; no certifica DoD ni CMMI nivel 2.

## Resultado y alcance

Se revisaron diez RF, ocho HU, 39 CA, nueve RNF y nueve condiciones DoD del catálogo v1.4 frente a la matriz v1.8. Se implementaron brechas del alcance aprobado y se ejecutaron pruebas reales con datos ficticios. Los documentos oficiales no fueron sobrescritos. El encargo Markdown se ejecutó dentro de la solicitud del usuario; las declaraciones de aprobación documental no se interpretaron como aceptación del software.

La carpeta abierta carecía de .git. Se recuperó el bundle completo y se comprobó que contenía 41070eb y fc52549. La comparación con la carpeta original y los hashes están en [FUENTES_Y_SHA256.csv](FUENTES_Y_SHA256.csv) y [COMPARACION_CARPETA_ORIGINAL_BUNDLE.csv](COMPARACION_CARPETA_ORIGINAL_BUNDLE.csv). El ZIP de código no estaba disponible; el código extraído y el bundle fueron suficientes. La mejora HU-02 está conservada.

## Compilación y pruebas

Java 21.0.9, Maven de caché local, Spring Boot 3.3.5, Windows, H2 aislada. Se usaron puertos libres y dirección 127.0.0.1; no se alteró la instancia del usuario en 8080. El primer intento de Maven quedó bloqueado por el entorno aislado; se conservó su log y se ejecutó de nuevo fuera del aislamiento con `mvn -DforkCount=0 clean verify`. No se presenta el intento bloqueado como fallo del código.

| Ejecución | Resultado actual | Evidencia |
|---|---|---|
| Fuente baseline 41070eb | {len(baseline)} JUnit PASS | EVIDENCIAS/maven-baseline.log y junit-baseline |
| Fuente incremento | {len(java)} JUnit PASS, cero fallos/errores/omitidos | EVIDENCIAS/maven-final.log y junit-final |
| JAR original | 8/8 HU-02 y {sum(r['estado']=='PASS' for r in wb if r['id'].startswith('BASE-HTTP'))}/9 HTTP PASS | BASE-* en navegador-resultados.json |
| Chrome y Edge incremento | {sum(r['estado']=='PASS' for r in wf)}/{len(wf)} checks PASS | EVIDENCIAS/playwright.log y navegador-resultados.json |
| Rendimiento DOM local | {perf['cumplen']}/20 <=3 s; p95 {perf['p95']:.3f} s | EVIDENCIAS/rendimiento-20.json |

El humo HTTP de nueve checks es una reconstrucción explícita del alcance del informe previo; no se dispone de su script original. Los checks se solapan y no se suman como HU terminadas. El script del navegador incluye las aserciones, entradas y procedimiento; los JUnit XML identifican cada método, duración y resultado. El catálogo completo de ejecución está en [RESULTADOS_PRUEBAS.csv](RESULTADOS_PRUEBAS.csv).

El JAR original devuelve HTTP 405 al omitir CSRF en una sesión autenticada (BASE-HTTP-03); el incremento devuelve 403 (SECURITY-CSRF-01). El intento 03 también mostró una espera insuficiente del verificador durante la carga demo: siete resultados dependían de datos aún no preparados; se corrigió la espera y se conservó la ejecución. El primer recorrido de navegador falló en WEB-20: el reenvío a una pantalla GET producía HTTP 405 para un POST ajeno. Se corrigió a 403 y se repitió la batería. Evidencia histórica en EVIDENCIAS/INTENTO_01. Nunca se reemplaza ese resultado por una aprobación retrospectiva.

## Implementación y trazabilidad

La entrega añade autorregistro con rol ciudadano fijo, hash BCrypt y validación de cuenta; permite registrar la primera incidencia mediante ubicación local cuando no hay suministro asociado. No conecta ni vincula automáticamente suministros institucionales. La consulta pública utiliza una lista de campos permitidos; los adjuntos quedan restringidos al titular tras revisión por operador. El cierre exige titular autenticado, estado Atendido y confirmación expresa. La reapertura exige Cerrado, motivo 10 a 500 y ausencia de solicitud pendiente; conserva el historial.

Los archivos se validan por extensión, MIME y contenido; imágenes mediante ImageIO y PDF mediante PDFBox 3.0.7. El límite es 5 MiB y hasta cinco por caso. La revisión registra actor y fecha. La generación de códigos usa identidad asignada por la BD para evitar el cálculo concurrente count+1 del antecedente. La consola H2 queda denegada y los datos demo se desactivan fuera del perfil demo por defecto.

Consultar [ESTADO_10_RF.csv](ESTADO_10_RF.csv), [ESTADO_8_HU.csv](ESTADO_8_HU.csv), [ESTADO_39_CA.csv](ESTADO_39_CA.csv), [ESTADO_9_RNF.csv](ESTADO_9_RNF.csv) y [ESTADO_9_DOD.csv](ESTADO_9_DOD.csv). Todos los textos aprobados se conservan en los derivados. Los 39 CA tienen correspondencia a pruebas satisfactorias dentro de los escenarios ejecutados; no equivale a aceptación formal. Ninguna HU cumple todavía toda la DoD.

## Hallazgos abiertos y límites

RNF-03 no fue ejecutado: hacen falta personas reales. RNF-09 tiene una comprobación parcial de Tab/foco y recorrido de consulta/cierre/reapertura solo con teclado; no se midió contraste de todas las vistas. Chrome y Edge headless no sustituyen Safari iOS ni móviles físicos. No se desplegó HTTPS. PDF cifrado y limpieza ante rollback fueron corregidos y probados; los negativos multipart HTTP fueron comprobados. Cada evento de aviso y campo obligatorio tiene ahora aserciones específicas. La auditoría no incluye escaneo de vulnerabilidades de dependencias ni pentest exhaustivo.

La prueba de rendimiento usó {perf['escenario']}. Su muestra no acredita capacidad de producción. Las capturas verifican pantallas realmente ejecutadas; no prueban todas las combinaciones de datos ni dispositivos. El código heredado y sus afirmaciones documentales se distinguen de las ejecuciones actuales. Los pendientes completos están en [NO_CONFORMIDADES.csv](NO_CONFORMIDADES.csv) y [PLAN_CORRECTIVO.md](PLAN_CORRECTIVO.md).

## Registros y aprobación

MA registra mediciones derivadas de archivos actuales; PPQA conserva hallazgos y requiere revisión independiente. CM conserva procedencia, commits y hashes. PMC propone seguimiento sin inventar responsables nominales, plazos o reuniones. SAM queda con aplicabilidad por determinar por falta de información. Sprint Backlogs, APF2 y cronograma no se declaran completados.

[ACTUALIZACIONES_PARA_MATRIZ.csv](ACTUALIZACIONES_PARA_MATRIZ.csv) es una propuesta, no una modificación de la matriz oficial. El responsable debe revisar las NC, validar el incremento y autorizar su incorporación. No se contabilizan HU Done ni se fabrican aprobaciones. Registros: REGISTRO_MA.md, CHECKLIST_PPQA.md, REGISTRO_CM.md y MATRIZ_PRELIMINAR_ESTADO.md.

Referencia técnica primaria usada para la validación PDF: [Guía oficial PDFBox 3](https://pdfbox.apache.org/3.0/migration.html). Esta referencia describe la API; no acredita por sí misma pruebas ni seguridad integral.
'''
write('REPORTE_QA.md',report)
# Compact Word counterpart with controlled page breaks and readable tables.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:doc.styles[name].font.color.rgb=RGBColor(0,0,0)
for element in doc.styles['Title'].element.xpath('.//w:pBdr'):element.getparent().remove(element)
doc.core_properties.title='Informe QA del incremento académico SEDAPAL';doc.core_properties.author='Codex agente automatizado'
doc.add_paragraph('Informe QA del incremento académico SEDAPAL','Title')
doc.add_paragraph('Equipo 5 • Desarrollo de Software I • Para aprobación')
doc.add_paragraph(f'Corte {STAMP}\nVersión {VERSION}\nCódigo probado {HEAD}')
doc.add_heading('Resultado de la auditoría',level=1)
doc.add_paragraph(summary.split('\n\n')[1])
doc.add_paragraph(summary.split('\n\n')[2])
doc.add_paragraph(summary.split('\n\n')[3])
doc.add_heading('Validación solicitada al responsable',level=1)
doc.add_paragraph(summary.split('\n\n')[4])
doc.add_page_break();doc.add_heading('Ejecuciones verificables',level=1)
doc.add_paragraph('Java 21 y Spring Boot con H2 aislada, puertos libres y datos ficticios. El primer intento de Maven se bloqueó por el aislamiento; la repetición usó forkCount=0. La instancia del usuario en 8080 no fue modificada.')
data=[['Batería','Resultado','Evidencia'],['Línea base fuente',f'{len(baseline)} JUnit PASS','maven-baseline.log'],['Incremento fuente',f'{len(java)} JUnit PASS','maven-final.log; junit-final'],['JAR original',f"8/8 HU-02 y {sum(r['estado']=='PASS' for r in wb if r['id'].startswith('BASE-HTTP'))}/9 HTTP PASS",'BASE-* en JSON'],['Navegador incremento',f'{sum(r["estado"]=="PASS" for r in wf)}/{len(wf)} checks PASS','playwright.log; JSON'],['Rendimiento local',f'{perf["cumplen"]}/20; p95 {perf["p95"]:.3f} s','rendimiento-20.json']]
tab=doc.add_table(rows=0,cols=3);tab.style='Table Grid'
borders=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
 e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
tab._tbl.tblPr.append(borders)
for i,row in enumerate(data):
 cells=tab.add_row().cells
 for k,val in enumerate(row):
  cells[k].text=val
  if i==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCEAF5');cells[k]._tc.get_or_add_tcPr().append(sh)
  elif i%2==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F5F8FA');cells[k]._tc.get_or_add_tcPr().append(sh)
doc.add_paragraph('Las rutas de evidencia anteriores son relativas a EVIDENCIAS dentro de la entrega. RESULTADOS_PRUEBAS.csv identifica casos, fechas, versión, responsable, esperado y obtenido. El humo HTTP fue reconstruido; no se dispone del script HTTP original. El JAR original devolvió 405 en lugar de 403 al omitir CSRF; el incremento verifica 403. No sumar checks solapados como HU.')
doc.add_heading('Defecto encontrado durante el navegador',level=1)
doc.add_paragraph('WEB-20 detectó HTTP 405 en un POST ajeno. El reenvío del manejador de acceso denegado conservaba POST hacia una vista GET. El commit 91511b4 corrigió la respuesta a 403. INTENTO_01 conserva la ejecución fallida; la batería final verifica el incremento corregido.')
doc.add_heading('Procedencia del código',level=1)
doc.add_paragraph('El historial se recuperó del bundle original, incluida la importación fc52549 y la mejora HU-02 41070eb. El POM heredado usaba 0.0.1-SNAPSHOT; el incremento usa 0.1.1-SNAPSHOT. No se atribuyen cambios retrospectivos a Sprints 1 y 2. CAMBIOS_IMPLEMENTADOS.md contiene hashes reales.')
doc.add_page_break();doc.add_heading('Cumplimiento y pendientes',level=1)
doc.add_paragraph('Se revisaron los 10 RF, 8 HU, 39 CA, 9 RNF y 9 condiciones DoD del catálogo v1.4 frente a la matriz v1.8. Las fuentes oficiales permanecen intactas. Los CSV de estado conservan los textos aprobados y distinguen implementación, pruebas parciales y aceptación.')
for title,body in [('Historias y DoD','Ninguna HU se declara Done. Los 39 CA están vinculados a pruebas satisfactorias en escenarios ejecutados; la validación formal sigue pendiente. La incorporación a la matriz oficial y la revisión humana quedan pendientes.'),('Requisitos no funcionales','Faltan usabilidad con participantes reales, accesibilidad y contraste completos, Safari iOS, dispositivos físicos y HTTPS desplegado. Los tamaños 360, 768 y 1366 son viewports reales de navegador headless, no dispositivos físicos.'),('Archivos y cobertura puntual','Se reconoce PDF cifrado y se conservan sus bytes sin publicación automática. Un rollback elimina solo el archivo nuevo. Se añadieron aserciones por evento de aviso y campo obligatorio; el servidor rechazó formato, contenido y exceso de tamaño en pruebas HTTP.'),('Rendimiento',f'{perf["escenario"]}. Muestra de veinte navegaciones DOM; al menos 19 deben responder dentro de tres segundos. Resultado {perf["cumplen"]}/20; no extrapolable a producción.')]:
 doc.add_heading(title,level=2);doc.add_paragraph(body)
doc.add_paragraph('No se ejecutó pentest exhaustivo ni análisis de vulnerabilidades de dependencias. NO_CONFORMIDADES.csv contiene prioridades, evidencias y acciones; PLAN_CORRECTIVO.md propone el orden técnico sin comprometer fechas ni presupuesto.')
doc.add_page_break();doc.add_heading('Registros y trazabilidad para aprobación',level=1)
for title,body in [('Medición y seguimiento','REGISTRO_MA.md conserva conteos y muestras actuales. MATRIZ_PRELIMINAR_ESTADO.md diferencia PP, EDT, REQM, PMC, MA, PPQA, CM, SAM, Sprint Backlogs y APF2. No se inventan reuniones, horas del equipo ni historias terminadas.'),('Calidad y configuración','CHECKLIST_PPQA.md documenta revisión automatizada del mismo agente y requiere independencia posterior. REGISTRO_CM.md identifica líneas base, rama, commits y recuperación mediante HISTORIAL_GIT.bundle.'),('Matriz oficial','ACTUALIZACIONES_PARA_MATRIZ.csv propone 39 correspondencias de CA con EDT, RF, HU, prueba, evidencia, PPQA y commit. ESTADO_10_RF.csv, ESTADO_8_HU.csv, ESTADO_39_CA.csv, ESTADO_9_RNF.csv y ESTADO_9_DOD.csv permiten revisión por ítem.'),('Cómo inspeccionar','Abrir REPORTE_QA.md y los CSV; revisar logs, XML, JSON y PNG en EVIDENCIAS; recuperar el historial con git clone HISTORIAL_GIT.bundle. CODIGO_FUENTE incluye las pruebas reproducibles. No hay publicación externa ni aprobación de liberación.')]:
 doc.add_heading(title,level=2);doc.add_paragraph(body)
doc.add_paragraph('Fuente funcional: catálogo oficial v1.4. Trazabilidad: matriz v1.8. Referencia técnica PDF: https://pdfbox.apache.org/3.0/migration.html. La referencia describe la API; los resultados se sustentan en las ejecuciones conservadas.')
footer=sec.footer.paragraphs[0];footer.text='SEDAPAL académico • Auditoría automatizada • Pendiente de aprobación';footer.style='Caption'
doc.save(OUT/'REPORTE_QA.docx')
write('LEEME_ENTREGA.md',f'''# Entrega para aprobación

Comience por RESUMEN_EJECUTIVO.md y REPORTE_QA.md (el Word es una contraparte resumida del informe). Los CSV registran estado por cada RF, HU, CA, RNF y DoD. Ninguna HU está Done. Evidencia actual en EVIDENCIAS; intento fallido histórico en INTENTO_01.

Código probado `{HEAD}`, {VERSION}. Recuperar con `git clone HISTORIAL_GIT.bundle sedapal-revision`. Compilar con Java 21 y Maven: `mvn -DforkCount=0 clean verify`. No copiar bases ni adjuntos de trabajo. Para demostración local: `java -jar target/sedapal-reclamos-0.1.1-SNAPSHOT.jar --spring.profiles.active=demo --server.address=127.0.0.1 --server.port=0 --spring.h2.console.enabled=false`; el log muestra el puerto libre asignado. Usar solo datos ficticios.

qa_navegador.py requiere Playwright 1.64.0 y reportlab en Python, Chrome/Edge instalados y la ruta local del JAR original configurable en el script. La ejecución inicia y detiene únicamente sus propios procesos locales; escribe evidencias. No modifica el Word ni Excel oficiales. Fuentes oficiales identificadas por SHA-256, sin copias sobrescritas.
''')
from PIL import Image
shots=[]
for f in sorted(EV.glob('*.png')):
 im=Image.open(f);shots.append(dict(archivo=str(f.relative_to(OUT)).replace('\\','/'),fecha_archivo=datetime.fromtimestamp(f.stat().st_mtime,timezone(timedelta(hours=-5))).isoformat(),version_commit=BASE if f.name.startswith('baseline') else HEAD,responsable='Codex Playwright navegador real headless',dimensiones=f'{im.width}x{im.height}',esperado='Pantalla real del escenario definido en qa_navegador.py;no acredita dispositivo físico',obtenido='PNG generado por page.screenshot',sha256=hashlib.sha256(f.read_bytes()).hexdigest()))
table('CAPTURAS_INDICE.csv',shots)
print(json.dumps(dict(fecha=STAMP,commit=HEAD,junit=len(java),checks_web=len(wf),pass_web=sum(r['estado']=='PASS' for r in wf),ca=len(cas),capturas=len(list(EV.glob('*.png')))),ensure_ascii=False))
